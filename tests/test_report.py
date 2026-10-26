"""Test de bout en bout : export d'exemple -> classeur de reporting."""
from openpyxl import load_workbook

import generate_sample
import report


def test_full_pipeline(tmp_path):
    src, out = tmp_path / "export.xlsx", tmp_path / "reporting.xlsx"
    generate_sample.write_xlsx(generate_sample.build_rows(seed=1), src)
    assert report.main([str(src), "-o", str(out), "--top", "5"]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["Accueil", "Synthèse", "CA mensuel", "Régions", "Top clients", "Top produits",
                             "Anomalies", "Données nettoyées"]
    assert len(wb["Synthèse"]._charts) == 3
    assert wb["Top clients"].max_row == 6                  # en-tête + top 5
    types = {c.value for c in wb["Anomalies"]["B"] if c.value}
    assert {"Doublon", "Ligne vide", "Date illisible ou impossible", "Montant HT illisible"} <= types


def test_home_sheet_guides_the_reader(tmp_path):
    src, out = tmp_path / "export.xlsx", tmp_path / "reporting.xlsx"
    generate_sample.write_xlsx(generate_sample.build_rows(seed=1), src, seed=1)
    assert report.main([str(src), "-o", str(out)]) == 0
    wb = load_workbook(out)
    assert wb.active.title == "Accueil"
    texts = " ".join(str(c.value) for row in wb["Accueil"].iter_rows() for c in row if c.value)
    assert "FICTIVES" in texts and "graine aléatoire 1" in texts      # origine des données
    assert "Colonnes attendues" in texts and "raison sociale" in texts  # adapter à ses données
    links = {c.hyperlink.location for row in wb["Accueil"].iter_rows() for c in row if c.hyperlink}
    assert "'Données nettoyées'!A1" in links
    for name in ("CA mensuel", "Régions", "Top clients", "Top produits", "Données nettoyées"):
        assert wb[name].freeze_panes == "A2"


def test_home_sheet_without_description(tmp_path):
    src, out = tmp_path / "export.xlsx", tmp_path / "reporting.xlsx"
    generate_sample.write_xlsx(generate_sample.build_rows(seed=1), src)  # pas de propriétés
    assert report.main([str(src), "-o", str(out)]) == 0
    texts = " ".join(str(c.value) for row in load_workbook(out)["Accueil"].iter_rows()
                     for c in row if c.value)
    assert "FICTIVES" not in texts and "Fichier source" in texts


def test_reference_month_option(tmp_path):
    src, out = tmp_path / "export.xlsx", tmp_path / "r.xlsx"
    generate_sample.write_xlsx(generate_sample.build_rows(seed=2), src)
    assert report.main([str(src), "-o", str(out), "--mois", "2025-06"]) == 0
    assert "06/2025" in load_workbook(out)["Synthèse"]["A2"].value


def test_missing_column_gives_clear_error(tmp_path, capsys):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["Date", "Client"])
    wb.active.append(["01/01/2025", "X"])
    src = tmp_path / "bad.xlsx"
    wb.save(src)
    assert report.main([str(src), "-o", str(tmp_path / "o.xlsx")]) == 1
    assert "Colonnes obligatoires introuvables" in capsys.readouterr().err


def test_missing_report_column_gives_clear_error(tmp_path, capsys):
    import pandas as pd
    src = tmp_path / "export.xlsx"
    generate_sample.write_xlsx(generate_sample.build_rows(seed=1), src)
    pd.read_excel(src, dtype=object).drop(columns=["Région"]).to_excel(src, index=False)
    assert report.main([str(src), "-o", str(tmp_path / "o.xlsx")]) == 1
    assert "region" in capsys.readouterr().err
