"""Écriture du classeur de reporting mis en forme (openpyxl, graphiques Excel natifs)."""
from __future__ import annotations

from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo

from .cleaning import COLUMNS, LABELS, REPORT_REQUIRED

# ---------------------------------------------------------------------------
# Charte graphique (à adapter aux couleurs du client)
# ---------------------------------------------------------------------------
# Palette sobre : un bleu marine, un bleu clair pour les fonds, des gris pour le texte.
PRIMARY = "1F4E79"      # titres, valeurs clés, séries des graphiques
LIGHT = "DDEBF7"        # fonds des indicateurs et des encadrés
GREY = "595959"         # textes secondaires
PALE_GREY = "F2F2F2"    # fonds neutres
TAB_MAIN, TAB_ANALYSIS, TAB_QUALITY = PRIMARY, "9DC3E6", "A6A6A6"  # couleurs d'onglets
TABLE_STYLE = "TableStyleMedium2"
EUR = '#,##0" €"'
EUR2 = '#,##0.00" €"'
PCT = '+0.0%;-0.0%;0.0%'
SHARE = "0.0%"
MONTH = "mmm yyyy"

TITLE = Font(size=18, bold=True, color=PRIMARY)
SUBTITLE = Font(size=10, italic=True, color=GREY)
H2 = Font(size=12, bold=True, color=PRIMARY)
KPI_LABEL = Font(size=10, color=GREY)
KPI_VALUE = Font(size=20, bold=True, color=PRIMARY)
KPI_FILL = PatternFill("solid", fgColor=LIGHT)
THIN = Side(style="thin", color="BFBFBF")
GREEN = Font(color="2E7D32")
RED = Font(color="C62828")
HEADER_FILL = PatternFill("solid", fgColor=PRIMARY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
BOX_FILL = PatternFill("solid", fgColor=LIGHT)
NEUTRAL_FILL = PatternFill("solid", fgColor=PALE_GREY)
LINK = Font(color="2E75B6", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")


def _autosize(ws, min_width=10, max_width=45):
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                n = len(c.value.strftime("%d/%m/%Y")) if isinstance(c.value, datetime) else len(str(c.value))
                widths[c.column_letter] = max(widths.get(c.column_letter, 0), n)
    for col, w in widths.items():
        ws.column_dimensions[col].width = max(min_width, min(max_width, w + 3))


def _write_table(ws, df: pd.DataFrame, name: str, top: int = 1, left: int = 1, formats=None,
                 variation_cols=()):
    """Écrit un DataFrame sous forme de Tableau Excel (filtres, lignes alternées)."""
    formats = formats or {}
    for j, col in enumerate(df.columns):
        ws.cell(top, left + j, col)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            if v is not None and not (isinstance(v, float) and pd.isna(v)):
                if isinstance(v, pd.Timestamp):
                    v = v.to_pydatetime()
                ws.cell(top + i, left + j, v)
    last_row = top + max(len(df), 1)
    last_col = left + len(df.columns) - 1
    ref = f"{get_column_letter(left)}{top}:{get_column_letter(last_col)}{last_row}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=TABLE_STYLE, showRowStripes=True)
    ws.add_table(t)
    for j, col in enumerate(df.columns):
        letter = get_column_letter(left + j)
        if col in formats:
            for r in range(top + 1, last_row + 1):
                ws[f"{letter}{r}"].number_format = formats[col]
        if col in variation_cols and len(df):
            rng = f"{letter}{top + 1}:{letter}{last_row}"
            ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=RED))
            ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], font=GREEN))
    return ref


def _bar(title, ws, cats_col, vals_col, top, n, horizontal=True, y_title="CA HT (€)", labels=True):
    ch = BarChart()
    ch.height, ch.width = max(7.5, 1.6 + 0.75 * n), 16
    ch.type = "bar" if horizontal else "col"
    ch.title = title
    ch.style = 2
    ch.legend = None
    ch.y_axis.title = y_title
    ch.y_axis.numFmt = '#,##0'
    ch.y_axis.majorGridlines = None
    data = Reference(ws, min_col=vals_col, min_row=top, max_row=top + n)
    cats = Reference(ws, min_col=cats_col, min_row=top + 1, max_row=top + n)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"  # le plus gros en haut
    ch.series[0].graphicalProperties.solidFill = PRIMARY
    ch.series[0].graphicalProperties.line.solidFill = PRIMARY
    ch.gapWidth = 60
    if labels:  # valeur affichée au bout de chaque barre
        ch.series[0].dLbls = DataLabelList(showVal=True)
        ch.series[0].dLbls.numFmt = '#,##0" €"'
        for attr in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
            setattr(ch.series[0].dLbls, attr, False)
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    return ch


def _line(ws, n, top=1):
    ch = LineChart()
    ch.title = "Chiffre d'affaires HT mensuel"
    ch.style = 2
    ch.legend = None
    ch.height, ch.width = 7.5, 22
    ch.y_axis.title = "CA HT (€)"
    ch.y_axis.numFmt = '#,##0'
    ch.x_axis.number_format = "mmm yy"
    ch.x_axis.majorTimeUnit = "months"
    data = Reference(ws, min_col=2, min_row=top, max_row=top + n)
    cats = Reference(ws, min_col=1, min_row=top + 1, max_row=top + n)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    s = ch.series[0]
    s.graphicalProperties.line.solidFill = PRIMARY
    s.graphicalProperties.line.width = 28000
    s.smooth = False
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    return ch


def _link(cell, sheet: str):
    """Lien interne vers la cellule A1 d'un onglet (pas de fichier externe)."""
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A1", display=str(cell.value))
    cell.font = LINK


def _home_link(ws, cell):
    """Lien « ← Accueil » en haut de chaque onglet."""
    ws[cell] = "← Retour à l'accueil"
    _link(ws[cell], "Accueil")


def _print_setup(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _kpi(ws, row, col, label, value, fmt, width=2):
    if value is not None and pd.isna(value):
        value = None
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + width - 1)
    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + width - 1)
    a, b = ws.cell(row, col, label), ws.cell(row + 1, col, value if value is not None else "n.d.")
    a.font, b.font = KPI_LABEL, KPI_VALUE
    b.number_format = fmt
    for r in (row, row + 1):
        for c in range(col, col + width):
            cell = ws.cell(r, c)
            cell.fill = KPI_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row + 1].height = 32
    if fmt == PCT and value is not None:
        b.font = Font(size=20, bold=True, color=GREEN.color if value >= 0 else RED.color)


# Exemples affichés dans l'onglet Accueil, colonne par colonne
EXAMPLES = {
    "date": "15/03/2025, 2025-03-15 ou vraie date Excel", "facture": "F24001",
    "client": "Boulangerie Martin", "produit": "Ramette papier A4", "region": "Bretagne",
    "commercial": "Sophie Laurent", "quantite": "3", "montant_ht": "1 234,50 € ou 1234.5",
    "tva": "246,90", "montant_ttc": "1 481,40",
}

SHEET_GUIDE = [  # onglet, contenu, astuce
    ("Synthèse", "Indicateurs du mois de référence et 3 graphiques.",
     "La page à imprimer ou à envoyer. Évolutions en vert (hausse) ou rouge (baisse)."),
    ("CA mensuel", "CA HT et nombre de factures par mois, évolutions M-1 et N-1, courbe.",
     "Les mois sans vente apparaissent à 0 € pour ne pas fausser les évolutions."),
    ("Régions", "CA HT, part du CA et nombre de clients par région.", "Filtres dans les en-têtes."),
    ("Top clients", "Les premiers clients par CA HT, avec leur part du CA.",
     "Taille du top réglable (option --top)."),
    ("Top produits", "Les premiers produits par CA HT, quantités vendues.", ""),
    ("Anomalies", "Corrections automatiques et lignes écartées ou à vérifier.",
     "Chaque ligne indique son numéro dans le fichier d'origine."),
    ("Données nettoyées", "Toutes les lignes retenues, prêtes pour un tableau croisé dynamique.",
     "Tableau Excel filtrable ; « Ligne source » renvoie au fichier d'origine."),
]


def _section(ws, row, title):
    ws.cell(row, 2, title).font = H2
    for col in range(2, 6):
        ws.cell(row, col).border = Border(bottom=Side(style="medium", color=PRIMARY))
    return row + 1


def _header_row(ws, row, labels):
    for j, text in enumerate(labels, start=2):
        c = ws.cell(row, j, text)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)


def _text_rows(ws, row, rows, fill=None):
    for values in rows:
        for j, v in enumerate(values, start=2):
            c = ws.cell(row, j, v)
            c.alignment = WRAP
            c.border = Border(bottom=THIN)
            if fill:
                c.fill = fill
        row += 1
    return row


def _paragraph(ws, row, text, height=None, font=None, fill=None):
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    c = ws.cell(row, 2, text)
    c.alignment = WRAP
    if font:
        c.font = font
    if fill:
        for col in range(2, 6):
            ws.cell(row, col).fill = fill
    if height:
        ws.row_dimensions[row].height = height
    return row + 1


def _home_sheet(ws, clean_df, quality, k, source_name, company, source_note, sheet_label):
    """Onglet Accueil : à quoi sert le classeur, comment le lire, d'où viennent les données."""
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TAB_MAIN
    for col, w in zip("ABCDE", (3, 24, 52, 52, 3)):
        ws.column_dimensions[col].width = w

    ws["B1"], ws["B1"].font = f"Reporting des ventes — {company}", TITLE
    ws.row_dimensions[1].height = 30
    ws["B2"] = f"Généré le {datetime.now():%d/%m/%Y à %H:%M} à partir de « {source_name} »."
    ws["B2"].font = SUBTITLE

    # En bref
    row = _section(ws, 4, "En bref")
    first, last = clean_df["date"].min(), clean_df["date"].max()
    brief = [
        ("Période couverte", f"du {first:%d/%m/%Y} au {last:%d/%m/%Y}"),
        ("Mois de référence", f"{k.mois_ref:%m/%Y}"),
        ("CA HT du mois", f"{k.ca_mois:,.0f} €".replace(",", " ")),
        ("Lignes de l'export", f"{quality.rows_read} lues, {quality.rows_kept} retenues, "
                               f"{quality.n_anomalies} signalées dans « Anomalies »"),
    ]
    for label, value in brief:
        ws.cell(row, 2, label).font = KPI_LABEL
        c = ws.cell(row, 3, value)
        c.font = Font(bold=True, color=PRIMARY)
        for col in (2, 3, 4):
            ws.cell(row, col).fill = BOX_FILL
        row += 1

    # Mode d'emploi
    row = _section(ws, row + 1, "Comment lire ce classeur")
    _header_row(ws, row, ["Onglet", "Contenu", "Astuce"])
    row += 1
    for name, content, tip in SHEET_GUIDE:
        _text_rows(ws, row, [(name, content, tip)])
        _link(ws.cell(row, 2), name)
        row += 1
    row = _paragraph(ws, row, "Cliquer sur un nom d'onglet pour s'y rendre. Chaque onglet a un lien "
                               "« ← Retour à l'accueil » en haut. Montants en euros hors taxes (HT).",
                     font=SUBTITLE)

    # Origine des données
    row = _section(ws, row + 1, "D'où viennent les données")
    origin = [
        ("Fichier source", source_name),
        ("Onglet lu", sheet_label),
        ("Traitement", "Lecture de l'export brut, nettoyage (reporting/cleaning.py), calculs "
                       "(reporting/aggregates.py), mise en forme (reporting/excel_report.py). "
                       "Le fichier source n'est jamais modifié."),
    ]
    row = _text_rows(ws, row, origin)
    if source_note:
        row = _paragraph(ws, row, source_note, height=15 * (2 + len(source_note) // 95),
                         fill=BOX_FILL)
    corrections = "; ".join(f"{label.lower()} : {n}" for label, n in quality.corrections.items())
    types = quality.anomalies["Type"].value_counts() if len(quality.anomalies) else {}
    flagged = "; ".join(f"{t.lower()} : {n}" for t, n in types.items()) or "aucune"
    row = _paragraph(ws, row, f"Corrections automatiques sur ce fichier — {corrections or 'aucune'}.",
                     height=32)
    row = _paragraph(ws, row, f"Lignes écartées ou à vérifier — {flagged}. Détail ligne par ligne "
                              "dans l'onglet « Anomalies ».", height=32)

    # Mise à jour
    row = _section(ws, row + 1, "Mettre à jour le reporting")
    steps = [
        ("1. Exporter", "Exporter les ventes depuis le logiciel (ERP, caisse, facturation) au format .xlsx."),
        ("2. Lancer", "python report.py export.xlsx -o reporting.xlsx   (options : --mois AAAA-MM, "
                      "--top 15, --onglet \"Ventes\", --societe \"Ma PME\")"),
        ("3. Ouvrir", "Ouvrir le nouveau classeur. Ne pas retoucher celui-ci à la main : il est "
                      "entièrement régénéré à chaque exécution."),
    ]
    row = _text_rows(ws, row, steps)

    # Colonnes attendues
    row = _section(ws, row + 1, "Colonnes attendues dans l'export")
    _header_row(ws, row, ["Colonne", "Intitulés reconnus (casse et accents ignorés)", "Exemple de valeur"])
    row += 1
    for name, aliases in COLUMNS.items():
        label = LABELS.get(name, name) + (" (obligatoire)" if name in REPORT_REQUIRED else " (facultative)")
        row = _text_rows(ws, row, [(label, ", ".join(aliases), EXAMPLES.get(name, ""))],
                         fill=NEUTRAL_FILL if name in REPORT_REQUIRED else None)
    row = _paragraph(ws, row, "Autres intitulés : les ajouter dans COLUMNS (reporting/cleaning.py). "
                              "Les colonnes en plus sont ignorées ; l'ordre des colonnes est libre.",
                     font=SUBTITLE)
    _print_setup(ws)
    ws.page_setup.orientation = "portrait"


def build_workbook(clean_df, quality, monthly_df, region_df, clients_df, products_df, k,
                   source_name: str, company: str = "Société exemple", source_note: str = None,
                   sheet_label: str = "premier onglet") -> Workbook:
    wb = Workbook()
    home = wb.active
    home.title = "Accueil"
    syn = wb.create_sheet("Synthèse")

    # --- Onglets de détail (écrits d'abord : les graphiques de la synthèse s'y réfèrent) ---
    ws_m = wb.create_sheet("CA mensuel")
    m = monthly_df.rename(columns={"mois": "Mois", "ca_ht": "CA HT", "nb_factures": "Nb factures",
                                   "var_m1": "Évol. M-1", "var_n1": "Évol. N-1"})
    _write_table(ws_m, m, "CA_mensuel", formats={"Mois": MONTH, "CA HT": EUR, "Nb factures": "#,##0",
                                                  "Évol. M-1": PCT, "Évol. N-1": PCT},
                 variation_cols=("Évol. M-1", "Évol. N-1"))
    _autosize(ws_m)
    ws_m.add_chart(_line(ws_m, len(m)), "G3")
    _home_link(ws_m, "G1")

    ws_r = wb.create_sheet("Régions")
    r = region_df.rename(columns={"region": "Région", "ca_ht": "CA HT", "nb_clients": "Nb clients",
                                  "part": "Part du CA"})
    _write_table(ws_r, r, "Regions", formats={"CA HT": EUR, "Part du CA": SHARE})
    _autosize(ws_r)
    ws_r.add_chart(_bar("CA HT par région", ws_r, 1, 2, 1, len(r)), "F3")
    _home_link(ws_r, "F1")

    ws_c = wb.create_sheet("Top clients")
    c = clients_df.rename(columns={"client": "Client", "ca_ht": "CA HT", "nb_factures": "Nb factures",
                                   "part": "Part du CA"})
    _write_table(ws_c, c, "Top_clients", formats={"CA HT": EUR, "Part du CA": SHARE})
    _autosize(ws_c)
    ws_c.add_chart(_bar(f"Top {len(c)} clients (CA HT)", ws_c, 1, 2, 1, len(c)), "F3")
    _home_link(ws_c, "F1")

    ws_p = wb.create_sheet("Top produits")
    p = products_df.rename(columns={"produit": "Produit", "ca_ht": "CA HT", "quantite": "Quantité",
                                    "part": "Part du CA"})
    _write_table(ws_p, p, "Top_produits", formats={"CA HT": EUR, "Quantité": "#,##0", "Part du CA": SHARE})
    _autosize(ws_p)
    ws_p.add_chart(_bar(f"Top {len(p)} produits (CA HT)", ws_p, 1, 2, 1, len(p)), "F3")
    _home_link(ws_p, "F1")

    for ws in (ws_m, ws_r, ws_c, ws_p):
        ws.freeze_panes = "A2"
        ws.sheet_properties.tabColor = TAB_ANALYSIS
        _print_setup(ws)

    # --- Anomalies ---
    ws_a = wb.create_sheet("Anomalies")
    ws_a["A1"], ws_a["A1"].font = "Contrôle qualité de l'export", H2
    ws_a["A2"] = (f"Lignes lues : {quality.rows_read} — lignes retenues : {quality.rows_kept}. "
                  "Numéros de ligne = lignes du fichier d'origine (en-tête en ligne 1).")
    ws_a["A2"].font = SUBTITLE
    row = 4
    ws_a.cell(row, 1, "Corrections automatiques").font = HEADER_FONT
    ws_a.cell(row, 2, "Nombre").font = HEADER_FONT
    ws_a.cell(row, 1).fill = ws_a.cell(row, 2).fill = HEADER_FILL
    for label, n in quality.corrections.items():
        row += 1
        ws_a.cell(row, 1, label)
        ws_a.cell(row, 2, n).number_format = "#,##0"
    row += 2
    ws_a.cell(row, 1, "Lignes à vérifier ou exclues").font = H2
    if len(quality.anomalies):
        _write_table(ws_a, quality.anomalies, "Anomalies", top=row + 1)
        ws_a.freeze_panes = f"A{row + 2}"
    else:
        ws_a.cell(row + 1, 1, "Aucune anomalie détectée.")
    _autosize(ws_a, max_width=50)
    ws_a.column_dimensions["A"].width = 34
    _home_link(ws_a, "G1")
    ws_a.sheet_properties.tabColor = TAB_QUALITY
    _print_setup(ws_a)

    # --- Données nettoyées ---
    ws_d = wb.create_sheet("Données nettoyées")
    cols = [c for c in ["ligne", "date", "mois", "facture", "client", "produit", "region", "commercial",
                        "quantite", "montant_ht", "tva", "montant_ttc"] if c in clean_df]
    d = clean_df[cols].rename(columns={**LABELS, "ligne": "Ligne source"})
    _write_table(ws_d, d, "Donnees", formats={"Date": "dd/mm/yyyy", "Mois": MONTH, "Quantité": "#,##0",
                                              "Montant HT": EUR2, "TVA": EUR2, "Montant TTC": EUR2})
    ws_d.freeze_panes = "A2"
    _autosize(ws_d)
    _home_link(ws_d, f"{get_column_letter(len(cols) + 2)}1")
    ws_d.sheet_properties.tabColor = TAB_QUALITY

    # --- Synthèse ---
    syn.sheet_view.showGridLines = False
    for col in range(1, 13):
        syn.column_dimensions[get_column_letter(col)].width = 13
    syn["A1"], syn["A1"].font = f"Reporting des ventes — {company}", TITLE
    syn.row_dimensions[1].height = 30
    syn["A2"] = (f"Mois de référence : {k.mois_ref:%m/%Y}  •  Source : {source_name}  •  "
                 f"Généré le {datetime.now():%d/%m/%Y à %H:%M}")
    syn["A2"].font = SUBTITLE
    _home_link(syn, "K2")
    _kpi(syn, 4, 1, f"CA HT {k.mois_ref:%m/%Y}", k.ca_mois, EUR)
    _kpi(syn, 4, 3, "Évolution vs mois précédent", k.var_m1, PCT)
    _kpi(syn, 4, 5, "Évolution vs même mois N-1", k.var_n1, PCT)
    _kpi(syn, 4, 7, f"CA HT cumulé {k.mois_ref.year}", k.ca_cumul, EUR)
    _kpi(syn, 4, 9, f"Évol. cumul vs {k.mois_ref.year - 1}", k.var_cumul, PCT)
    _kpi(syn, 4, 11, "Panier moyen du mois", k.panier_moyen_mois, EUR)
    _kpi(syn, 7, 1, "Factures du mois", k.nb_factures_mois, "0")
    _kpi(syn, 7, 3, "Clients actifs (12 mois)", k.clients_actifs_12m, "0")
    _kpi(syn, 7, 5, "Lignes à vérifier", quality.n_anomalies, "0")
    note = syn.cell(7, 7, "→ détail dans l'onglet « Anomalies »")
    _link(note, "Anomalies")
    note.alignment = Alignment(vertical="center")
    syn.merge_cells("G7:J8")

    syn.add_chart(_line(ws_m, len(m)), "A11")
    ch_r = _bar("CA HT par région", ws_r, 1, 2, 1, len(r))
    ch_r.height, ch_r.width = 8, 11
    syn.add_chart(ch_r, "A27")
    ch_c = _bar("Top 5 clients", ws_c, 1, 2, 1, min(5, len(c)))
    ch_c.height, ch_c.width = 8, 11
    syn.add_chart(ch_c, "G27")
    _print_setup(syn)
    syn.page_setup.fitToHeight = 1  # la synthèse tient sur une page
    syn.sheet_properties.tabColor = TAB_MAIN

    _home_sheet(home, clean_df, quality, k, source_name, company, source_note, sheet_label)
    wb.active = 0
    return wb
